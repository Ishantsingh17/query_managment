"""UC-01's compiled tabular evidence — the GL transaction listing.

Not every required item is a document sitting in a source system. This one is
the underlying row-level data, gathered from every source through MCP and
compiled into a single Excel workbook that then travels through validation and
packaging like any retrieved document.
"""

from __future__ import annotations

import sqlite3

UC01_QUERY = "Provide AP Cost Drill for August 2026, SOB 101, NAC 5000-5999."
UC01_QUERY_AR = "Provide AR Cost Drill for August 2026, SOB 101, NAC 5000-5999."

HEADER_ROW = 6
SOB_COLUMN = 4
NAC_COLUMN = 5
REPORT_TYPE_COLUMN = 7


def _run(client, query: str = UC01_QUERY) -> str:
    created = client.post("/api/audit-requests", json={"query": query})
    request_id = created.json()["request_id"]
    client.post(f"/api/audit-requests/{request_id}/run")
    return request_id


def _generated(request_id: str) -> dict:
    from app.services import repository

    return next(
        item
        for item in repository.list_retrieved_evidence(request_id)
        if item["document_type"] == "GL_TRANSACTION_LISTING"
    )


def _in_scope_row_count(report_type: str = "AP") -> int:
    """Count the rows the extract ought to contain, straight from the sources.

    Mirrors the applied criteria: period, SOB, NAC range and report type. Rows
    with no report_type are excluded, since they cannot be positively
    attributed to the requested type.
    """
    from app.catalog.loader import enabled_databases
    from app.config import get_settings

    settings = get_settings()
    expected = 0
    for spec in enabled_databases():
        connection = sqlite3.connect(spec.resolved_path(settings.data_root))
        connection.row_factory = sqlite3.Row
        try:
            for raw in connection.execute("SELECT * FROM transactions"):
                row = dict(raw)
                nac = row["nac_code"]
                if (
                    row["period"] == "August 2026"
                    and row["sob"] == "101"
                    and row["report_type"] == report_type
                    and nac
                    and nac.isdigit()
                    and 5000 <= int(nac) <= 5999
                ):
                    expected += 1
        finally:
            connection.close()
    return expected


def test_the_listing_is_compiled_not_retrieved(client):
    """It is recorded as a compiled extract, with no single source system."""
    request_id = _run(client)
    detail = client.get(f"/api/audit-requests/{request_id}").json()

    row = next(
        item
        for item in detail["retrieved_evidence"]
        if item["document_type"] == "GL_TRANSACTION_LISTING"
    )
    assert row["source_database_id"] == "GENERATED"
    assert row["source_database_name"] == "Compiled extract"
    assert row["staged_file_path"].endswith(".xlsx")
    assert row["has_file"] is True

    metadata = _generated(request_id)["metadata"]
    assert metadata["generated"] == "tabular"
    # Every source contributed rows, so the extract really does aggregate.
    assert sorted(metadata["contributing_databases"]) == [
        "DB-01",
        "DB-02",
        "DB-03",
        "DB-04",
    ]
    assert metadata["row_count"] > 0
    assert metadata["total_amount"] > 0


def test_the_listing_is_filtered_to_the_requested_report_type(client):
    """An AP request yields AP lines only, and the filter is recorded."""
    from openpyxl import load_workbook

    from app.config import get_settings

    request_id = _run(client)
    generated = _generated(request_id)
    metadata = generated["metadata"]

    assert metadata["applied_filters"] == {"report_type": "AP"}
    assert metadata["row_count"] == _in_scope_row_count("AP")

    workbook = load_workbook(get_settings().data_root / generated["staged_file_path"])
    sheet = workbook.active
    types = {
        sheet.cell(row=offset, column=REPORT_TYPE_COLUMN).value
        for offset in range(HEADER_ROW + 1, HEADER_ROW + 1 + metadata["row_count"])
    }
    assert types == {"AP"}, f"the extract carries other report types: {types}"

    # The criteria block states the filter, so the file explains its own scope.
    assert "Report Type: AP" in sheet["A3"].value


def test_the_filter_follows_the_query(client):
    """Asking for AR yields AR lines, not AP."""
    from openpyxl import load_workbook

    from app.config import get_settings

    request_id = _run(client, UC01_QUERY_AR)
    generated = _generated(request_id)
    metadata = generated["metadata"]

    assert metadata["applied_filters"] == {"report_type": "AR"}
    assert metadata["row_count"] == _in_scope_row_count("AR")

    workbook = load_workbook(get_settings().data_root / generated["staged_file_path"])
    sheet = workbook.active
    types = {
        sheet.cell(row=offset, column=REPORT_TYPE_COLUMN).value
        for offset in range(HEADER_ROW + 1, HEADER_ROW + 1 + metadata["row_count"])
    }
    assert types == {"AR"}


def test_report_type_filtering_does_not_touch_the_document_checklist(client):
    """The AP, AR and Others reports are all still required and retrieved.

    The filter must apply to the compiled extract only. If it leaked into the
    shared matcher, an AP request would stop finding the AR and Others cost
    drill documents and the checklist could never complete.
    """
    request_id = _run(client)
    detail = client.get(f"/api/audit-requests/{request_id}").json()

    found = {item["document_type"] for item in detail["retrieved_evidence"]}
    assert {"COST_DRILL_AP", "COST_DRILL_AR", "COST_DRILL_OTHERS"} <= found
    assert detail["evidence_found_count"] == 6
    assert detail["validation"]["validation_status"] == "COMPLETE"


def test_the_workbook_enforces_the_requested_scope(client):
    """No other SOB, and every NAC code inside the requested range.

    This is the guard against a subtle failure: identifier enrichment during
    the document sweep discovers strong identifiers (a vendor id harvested
    from one row) and folding those into an aggregate would silently narrow
    the population - a full listing quietly becoming a subset of itself.
    """
    from openpyxl import load_workbook

    from app.config import get_settings

    request_id = _run(client)
    generated = _generated(request_id)
    row_count = generated["metadata"]["row_count"]

    workbook = load_workbook(get_settings().data_root / generated["staged_file_path"])
    sheet = workbook.active
    assert sheet.title == "Transactions"

    headers = [sheet.cell(row=HEADER_ROW, column=c).value for c in range(1, 15)]
    assert headers[0] == "Transaction ID"
    assert "Amount" in headers
    assert "Source System" in headers

    sobs: set[str] = set()
    nacs: list[int] = []
    for offset in range(HEADER_ROW + 1, HEADER_ROW + 1 + row_count):
        sobs.add(sheet.cell(row=offset, column=SOB_COLUMN).value)
        nacs.append(int(sheet.cell(row=offset, column=NAC_COLUMN).value))

    assert len(nacs) == row_count
    assert sobs == {"101"}, f"the extract pulled in another SOB: {sobs}"
    assert min(nacs) >= 5000 and max(nacs) <= 5999, "NAC codes outside the range"


def test_row_fetches_do_not_inflate_document_match_counts(client):
    """Row fetches are not document searches and must not be counted as such."""
    request_id = _run(client)
    detail = client.get(f"/api/audit-requests/{request_id}").json()

    # Five documents across the sources; the fetched rows must not leak in.
    total = sum(attempt["result_count"] for attempt in detail["database_attempts"])
    assert total == 5, f"document match counts distorted by row fetches: {total}"

    for step in detail["timeline"]:
        if step["key"].startswith("database_") and step["detail"]:
            first = step["detail"].split()[0]
            if first.isdigit():
                assert int(first) <= 5, step


def test_the_listing_is_served_as_a_spreadsheet(client):
    """View on the extract downloads a real xlsx."""
    request_id = _run(client)
    detail = client.get(f"/api/audit-requests/{request_id}").json()

    row = next(
        item
        for item in detail["retrieved_evidence"]
        if item["document_type"] == "GL_TRANSACTION_LISTING"
    )
    response = client.get(
        f"/api/audit-requests/{request_id}/evidence/{row['evidence_id']}/file"
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    # xlsx files are zip archives.
    assert response.content[:2] == b"PK"


def test_the_listing_ships_in_the_final_package(client):
    """The workbook is packaged alongside the PDFs."""
    from app.services import packaging

    request_id = _run(client)
    evidence_dir = packaging.package_dir(request_id) / "evidence"
    names = sorted(path.name for path in evidence_dir.iterdir())

    assert any(name.endswith(".xlsx") for name in names), names
    assert len([name for name in names if name.endswith(".pdf")]) == 5


def test_other_use_cases_are_unaffected(client):
    """Only UC-01 declares generated evidence; UC-04 is untouched."""
    created = client.post(
        "/api/audit-requests",
        json={"query": "Provide alternate testing documents for ABC Ltd, Invoice INV-12345."},
    )
    request_id = created.json()["request_id"]
    client.post(f"/api/audit-requests/{request_id}/run")
    detail = client.get(f"/api/audit-requests/{request_id}").json()

    assert detail["evidence_required_count"] == 5
    assert detail["evidence_found_count"] == 5
    assert detail["retry_count"] == 1
    assert all(
        item["source_database_id"] != "GENERATED"
        for item in detail["retrieved_evidence"]
    )
