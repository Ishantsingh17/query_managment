"""Tabular evidence.

Some required evidence is not a document sitting in a source system: it is the
underlying row-level data. A GL transaction listing might be a hundred lines
pulled from several systems, and what the auditor needs is one workbook they
can filter and tie back to the reports.

This module compiles those rows into an .xlsx and writes it into the request's
staging area, so from that point on it travels through validation and
packaging exactly like any retrieved document.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.services import staging

# (header, row accessor key, number format, column width)
_COLUMNS: list[tuple[str, str, str | None, int]] = [
    ("Transaction ID", "transaction_id", None, 24),
    ("Posted Date", "posted_date", None, 18),
    ("Period", "period", None, 16),
    ("SOB", "sob", None, 8),
    ("NAC Code", "nac_code", None, 11),
    ("Account", "account_number", None, 11),
    ("Report Type", "report_type", None, 13),
    ("Vendor ID", "vendor_id", None, 11),
    ("Vendor Name", "vendor_name", None, 24),
    ("Description", "description", None, 34),
    ("Cost Centre", "cost_centre", None, 13),
    ("Amount", "amount", "#,##0.00", 14),
    ("Currency", "currency", None, 10),
    ("Source System", "source_system", None, 16),
]

_HEADER_FILL = PatternFill("solid", fgColor="0F1C3F")
_TOTAL_FILL = PatternFill("solid", fgColor="E2E8F0")


def _flatten(row: dict[str, Any], database_names: dict[str, str]) -> dict[str, Any]:
    """One transaction row in the shape the sheet expects."""
    attributes = row.get("attributes") or {}
    database_id = row.get("source_database_id") or ""
    return {
        "transaction_id": row.get("transaction_id"),
        "posted_date": attributes.get("posted_date"),
        "period": row.get("period"),
        "sob": row.get("sob"),
        "nac_code": row.get("nac_code"),
        "account_number": row.get("account_number"),
        "report_type": row.get("report_type"),
        "vendor_id": row.get("vendor_id"),
        "vendor_name": row.get("vendor_name"),
        "description": attributes.get("description"),
        "cost_centre": attributes.get("cost_centre"),
        "amount": row.get("amount"),
        "currency": row.get("currency"),
        "source_system": database_names.get(database_id, database_id),
    }


def build_transaction_workbook(
    request_id: str,
    *,
    label: str,
    filename: str,
    rows: list[dict[str, Any]],
    search_parameters: dict[str, Any],
    database_names: dict[str, str],
    parameter_labels: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Write the rows to an .xlsx in staging and describe what was written.

    Returns the staged relative path, row count, total value and the systems
    that contributed - all of which end up in the evidence metadata so the
    package can be audited without opening the file.
    """
    paths = staging.ensure_staging(request_id)
    destination: Path = paths["retrieved"] / filename

    flattened = [_flatten(row, database_names) for row in rows]
    # Stable, auditor-friendly ordering.
    flattened.sort(
        key=lambda item: (
            str(item.get("period") or ""),
            str(item.get("nac_code") or ""),
            str(item.get("transaction_id") or ""),
        )
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"

    # --- criteria block, so the extract states its own scope -------------
    labels = parameter_labels or {}
    sheet["A1"] = label
    sheet["A1"].font = Font(bold=True, size=13)
    sheet["A2"] = f"Request {request_id}"
    sheet["A2"].font = Font(size=10, color="65758C")

    criteria = [
        f"{labels.get(key, key.replace('_', ' ').title())}: {value}"
        for key, value in search_parameters.items()
        if value not in (None, "", [])
    ]
    sheet["A3"] = "Criteria — " + (", ".join(criteria) if criteria else "none")
    sheet["A3"].font = Font(size=10, color="65758C")
    # Registry order, so the list reads in the sequence the sources were
    # searched rather than alphabetically.
    contributed = [
        name
        for name in database_names.values()
        if any(item["source_system"] == name for item in flattened)
    ]
    sheet["A4"] = f"{len(flattened)} line(s) compiled from: " + ", ".join(contributed)
    sheet["A4"].font = Font(size=10, color="65758C")

    header_row = 6
    for index, (header, _key, _fmt, width) in enumerate(_COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width

    for offset, item in enumerate(flattened, start=header_row + 1):
        for index, (_header, key, number_format, _width) in enumerate(_COLUMNS, start=1):
            cell = sheet.cell(row=offset, column=index, value=item.get(key))
            if number_format:
                cell.number_format = number_format

    # --- total row -------------------------------------------------------
    total_value = sum(
        float(item["amount"]) for item in flattened if isinstance(item.get("amount"), (int, float))
    )
    if flattened:
        total_row = header_row + len(flattened) + 1
        amount_column = next(
            index for index, (_h, key, _f, _w) in enumerate(_COLUMNS, start=1) if key == "amount"
        )
        label_cell = sheet.cell(row=total_row, column=amount_column - 1, value="Total")
        label_cell.font = Font(bold=True)
        label_cell.fill = _TOTAL_FILL
        total_cell = sheet.cell(row=total_row, column=amount_column, value=round(total_value, 2))
        total_cell.font = Font(bold=True)
        total_cell.fill = _TOTAL_FILL
        total_cell.number_format = "#,##0.00"

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(_COLUMNS))}{header_row + len(flattened)}"
    )

    workbook.save(destination)

    data_root = get_settings().data_root
    return {
        "staged_file_path": str(destination.relative_to(data_root).as_posix()),
        "staged_filename": filename,
        "row_count": len(flattened),
        "total_amount": round(total_value, 2),
        "currency": next(
            (item["currency"] for item in flattened if item.get("currency")), None
        ),
        "contributing_databases": sorted(
            {row.get("source_database_id") for row in rows if row.get("source_database_id")}
        ),
    }
